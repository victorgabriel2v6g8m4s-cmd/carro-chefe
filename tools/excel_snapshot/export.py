#!/usr/bin/env python3
"""Export a deterministic, reviewable snapshot from the Carro Chefe XLSM.

This tool is repository tooling only. It never executes VBA. It reads the
workbook, exports cached cell values, formulas, structured tables and the VBA
source code, then writes deterministic text artifacts that can be reviewed by
humans and agents.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import os
import re
import shutil
import sys
import tempfile
import unicodedata
import zipfile
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path
from typing import Any, Iterable

TOOL_VERSION = "1.0.0"
REPO_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_SOURCE = Path("anexos/financeiro/carro chefe.xlsm")
DEFAULT_OUTPUT = Path("anexos/financeiro/snapshot")
BOOTSTRAP_FILE = "BOOTSTRAP_REQUIRED.json"
MANIFEST_FILE = "manifest.json"

DEFAULT_MAX_SOURCE_BYTES = 25 * 1024 * 1024
DEFAULT_MAX_ZIP_ENTRIES = 5_000
DEFAULT_MAX_UNCOMPRESSED_BYTES = 250 * 1024 * 1024
DEFAULT_MAX_ROWS = 50_000
DEFAULT_MAX_COLUMNS = 512


class SnapshotError(RuntimeError):
    """Expected validation/export failure with a user-facing message."""


@dataclass(frozen=True)
class PackageInfo:
    source_size: int
    source_sha256: str
    zip_entries: int
    zip_uncompressed_bytes: int
    vba_present: bool
    vba_project_size: int | None
    vba_project_sha256: str | None


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def stable_json_bytes(value: Any) -> bytes:
    text = json.dumps(value, ensure_ascii=False, indent=2, sort_keys=True, allow_nan=False)
    return (text + "\n").encode("utf-8")


def write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(stable_json_bytes(value))


def relative_posix(path: Path, root: Path) -> str:
    return path.relative_to(root).as_posix()


def resolve_repo_path(raw: str | Path) -> Path:
    candidate = Path(raw)
    resolved = candidate.resolve() if candidate.is_absolute() else (REPO_ROOT / candidate).resolve()
    try:
        resolved.relative_to(REPO_ROOT)
    except ValueError as exc:
        raise SnapshotError(f"Caminho fora do repositório não é permitido: {resolved}") from exc
    return resolved


def portable_slug(name: str, fallback: str) -> str:
    normalized = unicodedata.normalize("NFKD", name).encode("ascii", "ignore").decode("ascii")
    normalized = re.sub(r"[^A-Za-z0-9._-]+", "_", normalized).strip("._-")
    return normalized or fallback


def unique_filename(name: str, fallback: str, used: set[str], suffix: str = "") -> str:
    base = portable_slug(name, fallback)
    candidate = f"{base}{suffix}"
    counter = 2
    while candidate.casefold() in used:
        candidate = f"{base}_{counter}{suffix}"
        counter += 1
    used.add(candidate.casefold())
    return candidate


def json_value(value: Any) -> Any:
    if value is None or isinstance(value, (str, bool, int)):
        return value
    if isinstance(value, float):
        if not math.isfinite(value):
            return str(value)
        return value
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, bytes):
        return value.hex()
    return str(value)


def csv_value(value: Any) -> Any:
    value = json_value(value)
    return "" if value is None else value


def validate_xlsm_package(
    source: Path,
    *,
    max_source_bytes: int = DEFAULT_MAX_SOURCE_BYTES,
    max_zip_entries: int = DEFAULT_MAX_ZIP_ENTRIES,
    max_uncompressed_bytes: int = DEFAULT_MAX_UNCOMPRESSED_BYTES,
) -> PackageInfo:
    if source.suffix.lower() != ".xlsm":
        raise SnapshotError(f"A fonte deve ser .xlsm: {source}")
    if not source.is_file():
        raise SnapshotError(f"Planilha não encontrada: {source}")

    source_size = source.stat().st_size
    if source_size > max_source_bytes:
        raise SnapshotError(
            f"Planilha excede o limite de {max_source_bytes} bytes: {source_size} bytes."
        )
    if not zipfile.is_zipfile(source):
        raise SnapshotError("O arquivo informado não é um pacote XLSM/ZIP válido.")

    with zipfile.ZipFile(source) as archive:
        infos = archive.infolist()
        if len(infos) > max_zip_entries:
            raise SnapshotError(f"Pacote possui entradas demais: {len(infos)} > {max_zip_entries}.")

        total_uncompressed = 0
        names: set[str] = set()
        for info in infos:
            name = info.filename.replace("\\", "/")
            path = Path(name)
            if name.startswith("/") or ".." in path.parts:
                raise SnapshotError(f"Entrada ZIP suspeita rejeitada: {name}")
            if info.flag_bits & 0x1:
                raise SnapshotError(f"Entrada ZIP criptografada não é suportada: {name}")
            total_uncompressed += info.file_size
            if total_uncompressed > max_uncompressed_bytes:
                raise SnapshotError(
                    "Conteúdo descompactado excede o limite de segurança "
                    f"({max_uncompressed_bytes} bytes)."
                )
            names.add(name)

        required = {"[Content_Types].xml", "xl/workbook.xml"}
        missing = sorted(required - names)
        if missing:
            raise SnapshotError(f"Pacote XLSM incompleto; faltam: {', '.join(missing)}")

        vba_name = "xl/vbaProject.bin"
        if vba_name in names:
            vba_bytes = archive.read(vba_name)
            vba_size = len(vba_bytes)
            vba_sha = sha256_bytes(vba_bytes)
            vba_present = True
        else:
            vba_size = None
            vba_sha = None
            vba_present = False

    return PackageInfo(
        source_size=source_size,
        source_sha256=sha256_file(source),
        zip_entries=len(infos),
        zip_uncompressed_bytes=total_uncompressed,
        vba_present=vba_present,
        vba_project_size=vba_size,
        vba_project_sha256=vba_sha,
    )


def effective_bounds(ws: Any) -> tuple[int, int]:
    """Return the max row/column containing a value, ignoring formatting-only cells."""
    cells = getattr(ws, "_cells", {})
    max_row = 0
    max_col = 0
    for cell in cells.values():
        if cell.value is not None:
            max_row = max(max_row, int(cell.row))
            max_col = max(max_col, int(cell.column))
    return max_row, max_col


def enforce_sheet_limits(title: str, rows: int, columns: int, max_rows: int, max_columns: int) -> None:
    if rows > max_rows:
        raise SnapshotError(f"Aba {title!r} excede o limite de linhas: {rows} > {max_rows}.")
    if columns > max_columns:
        raise SnapshotError(f"Aba {title!r} excede o limite de colunas: {columns} > {max_columns}.")


def cached_or_formula(values_ws: Any, formula_ws: Any, row: int, column: int) -> Any:
    cached = values_ws.cell(row=row, column=column).value
    if cached is not None:
        return cached
    formula_cell = formula_ws.cell(row=row, column=column)
    if getattr(formula_cell, "data_type", None) == "f":
        return formula_cell.value
    return formula_cell.value


def write_sheet_csv(
    values_ws: Any,
    formula_ws: Any,
    target: Path,
    *,
    max_rows: int,
    max_columns: int,
) -> tuple[int, int]:
    rows_a, cols_a = effective_bounds(values_ws)
    rows_b, cols_b = effective_bounds(formula_ws)
    rows = max(rows_a, rows_b)
    columns = max(cols_a, cols_b)
    enforce_sheet_limits(formula_ws.title, rows, columns, max_rows, max_columns)

    target.parent.mkdir(parents=True, exist_ok=True)
    with target.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        if rows and columns:
            for row in range(1, rows + 1):
                writer.writerow(
                    [
                        csv_value(cached_or_formula(values_ws, formula_ws, row, column))
                        for column in range(1, columns + 1)
                    ]
                )
    return rows, columns


def workbook_properties(workbook: Any) -> dict[str, Any]:
    props = workbook.properties
    keys = (
        "title",
        "subject",
        "creator",
        "keywords",
        "description",
        "lastModifiedBy",
        "created",
        "modified",
        "category",
        "contentStatus",
        "identifier",
        "language",
        "version",
    )
    return {key: json_value(getattr(props, key, None)) for key in keys}


def defined_names(workbook: Any) -> list[dict[str, Any]]:
    output: list[dict[str, Any]] = []
    collection = workbook.defined_names
    try:
        items: Iterable[tuple[str, Any]] = collection.items()
    except AttributeError:
        items = ((item.name, item) for item in collection.definedName)

    for name, item in items:
        output.append(
            {
                "name": name,
                "value": getattr(item, "attr_text", None) or getattr(item, "value", None),
                "local_sheet_id": getattr(item, "localSheetId", None),
                "hidden": bool(getattr(item, "hidden", False)),
            }
        )
    return sorted(output, key=lambda item: (item["name"].casefold(), item["local_sheet_id"] or -1))


def extract_formulas(formula_wb: Any, values_wb: Any) -> list[dict[str, Any]]:
    result: list[dict[str, Any]] = []
    for formula_ws in formula_wb.worksheets:
        values_ws = values_wb[formula_ws.title]
        cells = getattr(formula_ws, "_cells", {})
        for cell in sorted(cells.values(), key=lambda current: (current.row, current.column)):
            if getattr(cell, "data_type", None) != "f":
                continue
            result.append(
                {
                    "sheet": formula_ws.title,
                    "cell": cell.coordinate,
                    "formula": str(cell.value),
                    "cached_value": json_value(values_ws[cell.coordinate].value),
                }
            )
    return result


def table_columns(table: Any, values_ws: Any, min_col: int, min_row: int, max_col: int) -> list[str]:
    table_column_objects = list(getattr(table, "tableColumns", []) or [])
    if table_column_objects and len(table_column_objects) == (max_col - min_col + 1):
        return [str(getattr(item, "name", "") or "") for item in table_column_objects]
    return [str(values_ws.cell(row=min_row, column=column).value or "") for column in range(min_col, max_col + 1)]


def export_tables(
    formula_wb: Any,
    values_wb: Any,
    target_root: Path,
    *,
    max_rows: int,
    max_columns: int,
) -> list[dict[str, Any]]:
    from openpyxl.utils.cell import range_boundaries

    index: list[dict[str, Any]] = []
    used: set[str] = set()
    table_dir = target_root / "tables"
    table_dir.mkdir(parents=True, exist_ok=True)

    for formula_ws in formula_wb.worksheets:
        values_ws = values_wb[formula_ws.title]
        for table_name in sorted(formula_ws.tables.keys(), key=str.casefold):
            table = formula_ws.tables[table_name]
            min_col, min_row, max_col, max_row = range_boundaries(table.ref)
            enforce_sheet_limits(formula_ws.title, max_row, max_col, max_rows, max_columns)
            columns = table_columns(table, values_ws, min_col, min_row, max_col)
            rows = [
                [json_value(cached_or_formula(values_ws, formula_ws, row, column)) for column in range(min_col, max_col + 1)]
                for row in range(min_row + 1, max_row + 1)
            ]
            filename = unique_filename(str(table_name), "table", used, ".json")
            relative_path = f"tables/{filename}"
            payload = {
                "name": str(table_name),
                "display_name": str(getattr(table, "displayName", table_name)),
                "sheet": formula_ws.title,
                "ref": table.ref,
                "columns": columns,
                "rows": rows,
                "totals_row_shown": bool(getattr(table, "totalsRowShown", False)),
            }
            write_json(target_root / relative_path, payload)
            index.append(
                {
                    "name": str(table_name),
                    "sheet": formula_ws.title,
                    "ref": table.ref,
                    "columns": columns,
                    "row_count": len(rows),
                    "path": relative_path,
                }
            )
    return sorted(index, key=lambda item: (item["name"].casefold(), item["sheet"].casefold()))


def decode_vba_text(value: Any) -> str:
    if not isinstance(value, bytes):
        return str(value)
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return value.decode(encoding)
        except UnicodeDecodeError:
            continue
    return value.decode("utf-8", errors="replace")


def decode_vba_source(value: Any) -> str:
    text = decode_vba_text(value)
    return text.replace("\r\n", "\n").replace("\r", "\n")


def extract_vba(source: Path, target_root: Path, package: PackageInfo) -> dict[str, Any]:
    vba_dir = target_root / "vba"
    modules_dir = vba_dir / "modules"
    modules_dir.mkdir(parents=True, exist_ok=True)

    modules: list[dict[str, Any]] = []
    if package.vba_present:
        try:
            from oletools.olevba import VBA_Parser
        except ImportError as exc:
            raise SnapshotError(
                "oletools não está instalado. Instale tools/excel_snapshot/requirements.txt."
            ) from exc

        parser = VBA_Parser(str(source))
        try:
            if parser.detect_vba_macros():
                used: set[str] = set()
                for _, stream_path, vba_filename, vba_code in parser.extract_macros():
                    original_name = decode_vba_text(vba_filename or "module.txt")
                    original_suffix = Path(original_name).suffix.lower()
                    suffix = original_suffix if original_suffix in {".bas", ".cls", ".frm"} else ".txt"
                    stem = Path(original_name).stem or "module"
                    filename = unique_filename(stem, "module", used, suffix)
                    code = decode_vba_source(vba_code)
                    if code and not code.endswith("\n"):
                        code += "\n"
                    path = modules_dir / filename
                    path.write_text(code, encoding="utf-8", newline="\n")
                    modules.append(
                        {
                            "name": stem,
                            "original_filename": original_name,
                            "stream_path": decode_vba_text(stream_path or ""),
                            "path": f"vba/modules/{filename}",
                            "sha256": sha256_bytes(code.encode("utf-8")),
                            "bytes": len(code.encode("utf-8")),
                        }
                    )
        finally:
            parser.close()

    index = {
        "present": package.vba_present,
        "vba_project_sha256": package.vba_project_sha256,
        "vba_project_bytes": package.vba_project_size,
        "module_count": len(modules),
        "modules": sorted(modules, key=lambda item: item["path"].casefold()),
        "execution": "never",
        "note": "Código VBA extraído estaticamente; nenhum macro é executado por esta ferramenta.",
    }
    write_json(vba_dir / "index.json", index)
    return index


def generated_file_index(root: Path, exclude: set[str] | None = None) -> list[dict[str, Any]]:
    excluded = exclude or set()
    rows: list[dict[str, Any]] = []
    for path in sorted((item for item in root.rglob("*") if item.is_file()), key=lambda item: relative_posix(item, root)):
        relative = relative_posix(path, root)
        if relative in excluded:
            continue
        rows.append({"path": relative, "bytes": path.stat().st_size, "sha256": sha256_file(path)})
    return rows


def generated_readme(source_rel: str, package: PackageInfo, sheet_count: int, table_count: int, formula_count: int, vba: dict[str, Any]) -> str:
    return f"""# Snapshot legível da planilha financeira\n\nEste diretório é **gerado automaticamente** por `tools/excel_snapshot/export.py`. Não edite seus arquivos manualmente.\n\nFonte binária: `{source_rel}`  \nSHA-256 da fonte: `{package.source_sha256}`  \nVersão do gerador: `{TOOL_VERSION}`\n\n## Conteúdo\n\n- `manifest.json`: integridade da fonte e dos artefatos gerados;\n- `workbook.json`: metadados, abas, nomes definidos e índice das tabelas;\n- `formulas.json`: fórmulas e valores em cache salvos pelo Excel;\n- `sheets/*.csv`: valores por aba para leitura e diff;\n- `tables/*.json`: tabelas estruturadas em formato estável;\n- `vba/index.json` e `vba/modules/*`: código-fonte VBA extraído estaticamente.\n\nResumo: {sheet_count} abas, {table_count} tabelas, {formula_count} fórmulas e {vba['module_count']} módulos VBA extraídos.\n\n## Governança\n\nO `.xlsm` continua sendo a fonte binária original desta representação. O snapshot não transforma valores legados, custos ou preços em parâmetros aprovados. O ERP continua destinado a ser a fonte transacional oficial conforme a governança do projeto.\n\nA ferramenta **não executa macros**, não recalcula a planilha com o motor do Excel e não é dependência de runtime do site, API ou ERP.\n"""


def close_workbook(workbook: Any) -> None:
    try:
        workbook.close()
    finally:
        archive = getattr(workbook, "vba_archive", None)
        if archive is not None:
            try:
                archive.close()
            except Exception:
                pass


def build_snapshot(
    source: Path,
    target_root: Path,
    *,
    max_rows: int,
    max_columns: int,
) -> dict[str, Any]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise SnapshotError(
            "openpyxl não está instalado. Instale tools/excel_snapshot/requirements.txt."
        ) from exc

    package = validate_xlsm_package(source)
    target_root.mkdir(parents=True, exist_ok=True)

    formula_wb = load_workbook(source, data_only=False, keep_vba=True, read_only=False)
    values_wb = load_workbook(source, data_only=True, keep_vba=True, read_only=False)
    try:
        sheet_dir = target_root / "sheets"
        sheet_dir.mkdir(parents=True, exist_ok=True)
        used_sheet_files: set[str] = set()
        sheets: list[dict[str, Any]] = []

        for formula_ws in formula_wb.worksheets:
            values_ws = values_wb[formula_ws.title]
            filename = unique_filename(formula_ws.title, "sheet", used_sheet_files, ".csv")
            rows, columns = write_sheet_csv(
                values_ws,
                formula_ws,
                sheet_dir / filename,
                max_rows=max_rows,
                max_columns=max_columns,
            )
            sheets.append(
                {
                    "name": formula_ws.title,
                    "state": formula_ws.sheet_state,
                    "rows": rows,
                    "columns": columns,
                    "path": f"sheets/{filename}",
                    "tables": sorted(formula_ws.tables.keys(), key=str.casefold),
                }
            )

        tables = export_tables(
            formula_wb,
            values_wb,
            target_root,
            max_rows=max_rows,
            max_columns=max_columns,
        )
        formulas = extract_formulas(formula_wb, values_wb)
        vba = extract_vba(source, target_root, package)

        workbook_payload = {
            "tool_version": TOOL_VERSION,
            "source": relative_posix(source, REPO_ROOT),
            "source_sha256": package.source_sha256,
            "properties": workbook_properties(formula_wb),
            "calculation": {
                "mode": getattr(getattr(formula_wb, "calculation", None), "calcMode", None),
                "full_calc_on_load": getattr(getattr(formula_wb, "calculation", None), "fullCalcOnLoad", None),
                "force_full_calc": getattr(getattr(formula_wb, "calculation", None), "forceFullCalc", None),
            },
            "defined_names": defined_names(formula_wb),
            "sheets": sheets,
            "tables": tables,
        }
        write_json(target_root / "workbook.json", workbook_payload)
        write_json(target_root / "formulas.json", {"count": len(formulas), "formulas": formulas})

        readme = generated_readme(
            relative_posix(source, REPO_ROOT),
            package,
            len(sheets),
            len(tables),
            len(formulas),
            vba,
        )
        (target_root / "README.md").write_text(readme, encoding="utf-8", newline="\n")

        file_index = generated_file_index(target_root, exclude={MANIFEST_FILE})
        manifest = {
            "schema_version": 1,
            "generator": {
                "path": "tools/excel_snapshot/export.py",
                "version": TOOL_VERSION,
                "deterministic": True,
            },
            "source": {
                "path": relative_posix(source, REPO_ROOT),
                "bytes": package.source_size,
                "sha256": package.source_sha256,
                "zip_entries": package.zip_entries,
                "zip_uncompressed_bytes": package.zip_uncompressed_bytes,
            },
            "workbook": {
                "sheet_count": len(sheets),
                "table_count": len(tables),
                "formula_count": len(formulas),
                "vba_present": package.vba_present,
                "vba_project_bytes": package.vba_project_size,
                "vba_project_sha256": package.vba_project_sha256,
                "vba_module_count": vba["module_count"],
            },
            "generated_files": file_index,
            "generated_at": None,
            "generated_at_note": "O timestamp é omitido intencionalmente para manter o snapshot determinístico; use o histórico Git.",
        }
        write_json(target_root / MANIFEST_FILE, manifest)
        return manifest
    finally:
        close_workbook(formula_wb)
        close_workbook(values_wb)


def directory_map(root: Path) -> dict[str, str]:
    return {
        relative_posix(path, root): sha256_file(path)
        for path in root.rglob("*")
        if path.is_file()
    }


def compare_snapshot(expected: Path, actual: Path) -> list[str]:
    expected_map = directory_map(expected)
    actual_map = directory_map(actual)
    messages: list[str] = []
    for missing in sorted(actual_map.keys() - expected_map.keys()):
        messages.append(f"faltando no repositório: {missing}")
    for stale in sorted(expected_map.keys() - actual_map.keys()):
        messages.append(f"arquivo extra/desatualizado: {stale}")
    for common in sorted(expected_map.keys() & actual_map.keys()):
        if expected_map[common] != actual_map[common]:
            messages.append(f"conteúdo divergente: {common}")
    return messages


def read_bootstrap_marker(output: Path) -> dict[str, Any] | None:
    path = output / BOOTSTRAP_FILE
    if not path.is_file():
        return None
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise SnapshotError(f"Marcador de bootstrap inválido: {path}") from exc


def output_is_managed(output: Path) -> bool:
    if not output.exists():
        return True
    if not output.is_dir():
        return False
    entries = [path.name for path in output.iterdir()]
    if not entries:
        return True
    return (output / MANIFEST_FILE).is_file() or (output / BOOTSTRAP_FILE).is_file()


def install_snapshot(staged: Path, output: Path) -> None:
    if not output_is_managed(output):
        raise SnapshotError(
            f"Destino contém arquivos não gerenciados e não será apagado automaticamente: {output}"
        )
    output.parent.mkdir(parents=True, exist_ok=True)
    backup = output.with_name(f".{output.name}.backup-{os.getpid()}")
    if backup.exists():
        shutil.rmtree(backup)
    try:
        if output.exists():
            output.replace(backup)
        staged.replace(output)
        if backup.exists():
            shutil.rmtree(backup)
    except Exception:
        if output.exists() and output != staged:
            shutil.rmtree(output, ignore_errors=True)
        if backup.exists():
            backup.replace(output)
        raise


def check_bootstrap(source: Path, output: Path, marker: dict[str, Any]) -> None:
    expected_path = marker.get("source_path")
    expected_sha = marker.get("source_sha256")
    source_rel = relative_posix(source, REPO_ROOT)
    actual_sha = sha256_file(source)
    if expected_path != source_rel:
        raise SnapshotError(
            f"Bootstrap aponta para {expected_path!r}, mas a fonte atual é {source_rel!r}."
        )
    if expected_sha != actual_sha:
        raise SnapshotError(
            "A planilha mudou antes da inicialização do snapshot. Gere e versione o snapshot agora; "
            f"esperado {expected_sha}, atual {actual_sha}."
        )


def run_check(source: Path, output: Path, *, allow_bootstrap: bool, max_rows: int, max_columns: int) -> None:
    with tempfile.TemporaryDirectory(prefix="carro-chefe-xlsm-check-") as temp_dir:
        staged = Path(temp_dir) / "snapshot"
        build_snapshot(source, staged, max_rows=max_rows, max_columns=max_columns)

        if (output / MANIFEST_FILE).is_file():
            if (output / BOOTSTRAP_FILE).exists():
                raise SnapshotError(
                    f"{BOOTSTRAP_FILE} deve ser removido quando o snapshot real existe."
                )
            differences = compare_snapshot(output, staged)
            if differences:
                detail = "\n  - ".join(differences[:50])
                remainder = len(differences) - 50
                suffix = f"\n  - ... e mais {remainder} diferença(s)" if remainder > 0 else ""
                raise SnapshotError(
                    "Snapshot desatualizado. Execute tools/excel_snapshot/export.py e versione as mudanças:\n"
                    f"  - {detail}{suffix}"
                )
            print("Snapshot XLSM está sincronizado e íntegro.")
            return

        marker = read_bootstrap_marker(output)
        if allow_bootstrap and marker is not None:
            check_bootstrap(source, output, marker)
            print(
                "Bootstrap validado: o snapshot ainda não foi versionado, mas a fonte continua no SHA-256 "
                f"esperado ({marker['source_sha256']}). O exportador foi executado integralmente em diretório temporário."
            )
            return

        raise SnapshotError(
            "Snapshot ainda não foi inicializado. Execute tools/excel_snapshot/export.py e versione "
            f"{relative_posix(output, REPO_ROOT)}."
        )


def parser() -> argparse.ArgumentParser:
    result = argparse.ArgumentParser(
        description="Exporta XLSM para CSV/JSON/VBA textual sem executar macros."
    )
    result.add_argument("--source", default=str(DEFAULT_SOURCE), help="XLSM fonte relativo ao repositório.")
    result.add_argument("--output", default=str(DEFAULT_OUTPUT), help="Diretório de snapshot relativo ao repositório.")
    result.add_argument("--check", action="store_true", help="Gera em temporário e falha se o snapshot versionado divergir.")
    result.add_argument(
        "--allow-bootstrap",
        action="store_true",
        help="Aceita temporariamente BOOTSTRAP_REQUIRED.json quando o primeiro snapshot ainda não foi versionado.",
    )
    result.add_argument("--max-rows", type=int, default=DEFAULT_MAX_ROWS, help="Limite de linhas por aba.")
    result.add_argument("--max-columns", type=int, default=DEFAULT_MAX_COLUMNS, help="Limite de colunas por aba.")
    return result


def main(argv: list[str] | None = None) -> int:
    args = parser().parse_args(argv)
    try:
        source = resolve_repo_path(args.source)
        output = resolve_repo_path(args.output)
        if source == output or output in source.parents:
            raise SnapshotError("O diretório de saída não pode conter nem substituir a planilha fonte.")
        if args.max_rows < 1 or args.max_columns < 1:
            raise SnapshotError("--max-rows e --max-columns devem ser maiores que zero.")
        if args.allow_bootstrap and not args.check:
            raise SnapshotError("--allow-bootstrap só pode ser usado junto com --check.")

        if args.check:
            run_check(
                source,
                output,
                allow_bootstrap=args.allow_bootstrap,
                max_rows=args.max_rows,
                max_columns=args.max_columns,
            )
            return 0

        output.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(prefix=".xlsm-snapshot-", dir=output.parent) as temp_dir:
            staged = Path(temp_dir) / "snapshot"
            manifest = build_snapshot(
                source,
                staged,
                max_rows=args.max_rows,
                max_columns=args.max_columns,
            )
            install_snapshot(staged, output)
        print(
            "Snapshot gerado: "
            f"{relative_posix(output, REPO_ROOT)} | "
            f"SHA-256 fonte {manifest['source']['sha256']} | "
            f"{manifest['workbook']['sheet_count']} abas | "
            f"{manifest['workbook']['table_count']} tabelas | "
            f"{manifest['workbook']['vba_module_count']} módulos VBA."
        )
        return 0
    except SnapshotError as exc:
        print(f"erro: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
