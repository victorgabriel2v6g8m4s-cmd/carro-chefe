from __future__ import annotations

import json
import shutil
import tempfile
import unittest
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.table import Table

from tools.excel_snapshot import export as snapshot


class SnapshotExporterTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = Path(tempfile.mkdtemp(prefix="xlsm-export-test-", dir=snapshot.REPO_ROOT))

    def tearDown(self) -> None:
        shutil.rmtree(self.temp, ignore_errors=True)

    def sample_workbook(self) -> Path:
        source = self.temp / "sample.xlsm"
        wb = Workbook()
        ws = wb.active
        ws.title = "Insumos"
        ws.append(["Item", "Qtd", "Preço"])
        ws.append(["Carne", 2, 10])
        ws.append(["Pão", 3, 5])
        ws["D1"] = "Total"
        ws["D2"] = "=B2*C2"
        ws.add_table(Table(displayName="insumos", ref="A1:C3"))
        wb.save(source)
        return source

    def test_snapshot_is_deterministic_and_keeps_formula(self) -> None:
        source = self.sample_workbook()
        first = self.temp / "first"
        second = self.temp / "second"
        snapshot.build_snapshot(source, first, max_rows=100, max_columns=20)
        snapshot.build_snapshot(source, second, max_rows=100, max_columns=20)
        self.assertEqual(snapshot.directory_map(first), snapshot.directory_map(second))
        formulas = json.loads((first / "formulas.json").read_text(encoding="utf-8"))
        self.assertEqual(formulas["count"], 1)
        self.assertEqual(formulas["formulas"][0]["formula"], "=B2*C2")
        table = json.loads((first / "tables" / "insumos.json").read_text(encoding="utf-8"))
        self.assertEqual(table["rows"][0], ["Carne", 2, 10])

    def test_rejects_zip_path_traversal(self) -> None:
        source = self.temp / "bad.xlsm"
        with zipfile.ZipFile(source, "w") as archive:
            archive.writestr("[Content_Types].xml", "<Types />")
            archive.writestr("xl/workbook.xml", "<workbook />")
            archive.writestr("../escape.txt", "no")
        with self.assertRaises(snapshot.SnapshotError):
            snapshot.validate_xlsm_package(source)

    def test_bootstrap_rejects_changed_source(self) -> None:
        source = self.sample_workbook()
        output = self.temp / "snapshot"
        output.mkdir()
        marker = {
            "source_path": snapshot.relative_posix(source, snapshot.REPO_ROOT),
            "source_sha256": "0" * 64,
        }
        with self.assertRaises(snapshot.SnapshotError):
            snapshot.check_bootstrap(source, output, marker)

    def test_unmanaged_output_is_not_deleted(self) -> None:
        output = self.temp / "snapshot"
        output.mkdir()
        (output / "keep.txt").write_text("manual", encoding="utf-8")
        self.assertFalse(snapshot.output_is_managed(output))


if __name__ == "__main__":
    unittest.main()
